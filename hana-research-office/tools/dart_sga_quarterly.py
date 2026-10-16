"""DART 판매비와관리비 주석 → 분기별 계정 모델 (고정비·변동비 분류)

상장사의 2019년 이후 분기·반기·사업보고서에서 연결(기본) 또는 별도 재무제표 주석의
「판매비와관리비」 표를 찾아 계정별 3개월·누적 금액을 모으고, 분기 모델 엑셀을 만듭니다.
  - 1Q = 1분기 누적, 2Q·3Q = 해당 보고서의 3개월 열, 4Q = 사업보고서 연간 − 3분기 누적 (엑셀 수식)
  - 계정 오른쪽 열에 고정비/변동비 분류, 고정비·변동비 합계, 공시 합계와의 Check 행
  - 금액 단위는 백만원으로 통일 (원문 단위는 「목록」 시트에 기록)

필요: 환경변수 DART_API_KEY (텔레그램 봇 .env의 키와 같은 것), 네트워크에서 opendart.fss.or.kr 허용
실행:
    python dart_sga_quarterly.py 066970 --since 20190101 --out lnf_sga.xlsx          # 엘앤에프, 연결
    python dart_sga_quarterly.py 020150 --out lotteemat_sga.xlsx                    # 롯데에너지머티리얼즈
    python dart_sga_quarterly.py 066970 --basis 별도
    python dart_sga_quarterly.py 066970 --xml-dir xml/   # --keep-xml로 받아둔 원문 재사용 (네트워크 불필요)
옵션: --env 경로(.env), --keep-xml 폴더(원문 XML 보관, 검증용)
"""
import argparse
import re
import sys
import time
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dart_periodic_extract import load_key, corp_code, periodic_reports, fetch_doc  # noqa: E402

import warnings  # noqa: E402
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning  # noqa: E402
warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)

UNIT_RE = re.compile(r"단위\s*[:：]?\s*(백만원|천원|억원|원)")
TO_MN = {"원": 1e-6, "천원": 1e-3, "백만원": 1.0, "억원": 100.0}          # → 백만원
SGA_RE = re.compile(r"판매비\s*(와|및)\s*관리비")
NOTE_RE = re.compile(r"^\(?\d+\)?[.)]\s*\S")                            # 주석 번호 제목 (예: 23. 판매비와관리비)
KNOWN = ("급여", "퇴직급여", "복리후생비", "감가상각비", "지급수수료", "운반비", "여비교통비", "세금과공과",
         "보험료", "경상개발비", "연구개발비", "수선비", "통신비", "소모품비", "광고선전비", "접대비", "임차료")
TOTAL = {"합계", "계", "총계", "판매비와관리비", "판매비와관리비합계", "판매비와관리비계", "판매비및관리비", "판매비및관리비합계"}
SYN = {"기업업무추진비": "접대비", "유형자산감가상각비": "감가상각비", "사용권자산감가상각비": "사용권자산상각비",
       "무형자산감가상각비": "무형자산상각비", "판매비와관리비급여": "급여"}
DISPLAY = {"접대비": "접대비(기업업무추진비)"}

# 고정비·변동비 분류 (계정명 일치 → 근거). 목록에 없는 계정은 고정비 기본값 + 노란 표시
VARIABLE = {
    "운반비": "출하량 연동", "판매수수료": "매출 연동", "수출제비용": "수출 물량 연동", "수출비용": "수출 물량 연동",
    "포장비": "출하량 연동", "견본비": "샘플 출하 연동", "하자보수비": "판매량 연동", "판매보증비": "판매량 연동",
    "판매촉진비": "매출 연동", "보관료": "재고·물량 연동", "창고료": "재고·물량 연동", "대손상각비": "매출채권 연동(비경상)",
}
FIXED = {
    "급여": "인건비", "종업원단기급여": "인건비", "퇴직급여": "인건비", "복리후생비": "인건비 부대", "주식보상비용": "인건비",
    "여비교통비": "관리 활동", "세금과공과": "재산세 등", "임차료": "계약 고정", "지급임차료": "계약 고정",
    "감가상각비": "자산 상각", "무형자산상각비": "자산 상각", "사용권자산상각비": "자산 상각", "수선비": "설비 유지",
    "보험료": "계약 고정", "접대비": "재량 지출", "광고선전비": "재량 지출", "경상개발비": "연구 인력·시험",
    "연구개발비": "연구 인력·시험", "교육훈련비": "재량 지출", "통신비": "관리 활동", "소모품비": "준고정(일시 구매로 변동 큼)",
    "전력비": "사무·연구 시설", "가스사용료": "사무·연구 시설", "수도사용료": "사무·연구 시설", "수도광열비": "사무·연구 시설",
    "도서인쇄비": "관리 활동", "지급수수료": "준고정(용역·IT·법률 등)", "차량유지비": "관리 활동", "잡비": "기타",
    "공해관리비": "환경 관리", "회의비": "관리 활동", "협회비": "계약 고정", "건물관리비": "계약 고정",
    "외주용역비": "준고정(용역)", "용역비": "준고정(용역)", "행사비": "재량 지출", "기밀비": "재량 지출",
}


def classify(acct):
    if acct in VARIABLE:
        return "변동비", VARIABLE[acct], False
    if acct in FIXED:
        return "고정비", FIXED[acct], False
    return "고정비", "기본값 — 확인 필요", True


def _txt(el):
    return re.sub(r"\s+", " ", el.get_text(" ", strip=True)).strip()


def norm_acct(lbl):
    s = re.sub(r"\s+", "", lbl)
    s = re.sub(r",?판관비$", "", s)                 # XBRL 뷰어 표기 「운반비, 판관비」
    s = re.sub(r"^[\(\[]?\d+[\)\].]?", "", s)        # 앞 번호
    s = re.sub(r"\(.*?\)", "", s)                    # (주1), (환입) 등
    s = s.strip("·-:：")
    return SYN.get(s, s)


def num(c):
    s = re.sub(r"[\s원]", "", c)
    if s in ("-", "－", "—", "–"):
        return 0.0
    m = re.fullmatch(r"(\()?(-)?([\d,]+(?:\.\d+)?)(\))?", s)
    if not m:
        return None
    v = float(m.group(3).replace(",", ""))
    return -v if (m.group(1) or m.group(2)) else v


def parse_sga(rows, unit):
    """판관비 표 → ({계정: (3개월, 누적)}, 공시합계(3개월, 누적), 단위) / 판관비 표가 아니면 None"""
    hdr, data = [], []
    for r in rows:
        if not data and not any(num(c) is not None for c in r[1:]):
            hdr.append(r)
        else:
            data.append(r)
    m = UNIT_RE.search(" ".join(" ".join(r) for r in hdr))
    unit = m.group(1) if m else unit
    i3 = ic = None
    for r in hdr:
        labs = [c for c in r if "3개월" in c or "누적" in c]
        if labs:
            i3 = next((i for i, c in enumerate(labs) if "3개월" in c), None)
            ic = next((i for i, c in enumerate(labs) if "누적" in c), None)
    if ic is None:
        ic = 0                                       # 1분기·사업보고서: 첫 숫자 열(당분기·당기)
    accts, total = {}, None
    for r in data:
        if not r or not r[0].strip():
            continue
        vals = [num(c) for c in r[1:]]
        cum = vals[ic] if ic < len(vals) else None
        m3 = vals[i3] if (i3 is not None and i3 < len(vals)) else cum
        a = norm_acct(r[0])
        if a in TOTAL:
            total = (m3, cum)
        elif a and (cum is not None or m3 is not None):
            p = accts.get(a, (0.0, 0.0))             # 같은 계정이 두 줄이면 합산 (예: 접대비·기업업무추진비)
            accts[a] = ((p[0] or 0) + (m3 or 0), (p[1] or 0) + (cum or 0))
    names = set(accts)
    if sum(any(k == n for n in names) for k in KNOWN) < 3 or any("원재료" in n or "재공품" in n for n in names):
        return None                                  # 비용의 성격별 분류 등 다른 표 제외
    return accts, total, unit


def extract_sga(xml, basis="연결"):
    """주석 구간(연결/별도)에서 「판매비와관리비」 제목 뒤 첫 판관비 표를 찾음. 제목 없이도 합계행이 판관비인 표는 인정"""
    soup = BeautifulSoup(xml, "html.parser")
    sec, flag, unit, ctx = None, False, None, ""
    for el in soup.find_all(["title", "p", "table"]):
        if el.find_parent("table") is not None:
            continue
        if el.name == "title":
            t = _txt(el)
            sec = "연결" if re.search(r"연결\s*재무\s*제표\s*주석", t) else "별도" if re.search(r"재무\s*제표\s*주석", t) else None
            flag = False
            continue
        if el.name == "p":
            t = _txt(el)
            mu = UNIT_RE.search(t)
            if mu:
                unit = mu.group(1)
            if len(t) <= 60 and SGA_RE.search(t):
                flag, ctx = True, t
            elif NOTE_RE.match(t) and len(t) <= 40:
                flag = False
            continue
        rows = [[_txt(td) for td in tr.find_all(["th", "td", "te", "tu"])] for tr in el.find_all("tr")]
        rows = [r for r in rows if any(r)]
        if len(rows) <= 2 and sum(len(r) for r in rows) <= 3:
            mu = UNIT_RE.search(" ".join(" ".join(r) for r in rows))
            if mu:
                unit = mu.group(1)
            continue
        if sec != basis or len(rows) < 5:
            continue
        res = parse_sga(rows, unit)
        if res and (flag or res[1] is not None and any(SGA_RE.search(r[0]) for r in rows if r)):
            return {"accts": res[0], "total": res[1], "unit": res[2] or "원", "context": ctx}
    return None


# ── 엑셀 ──────────────────────────────────────────────────────────────
KR, EN = "맑은 고딕", "Arial"
HAN = re.compile(r"[ᄀ-ᇿ㄰-㆏가-힣]")
LAT = re.compile(r"[A-Za-z0-9]")


def split_runs(s):
    kinds = [True if HAN.match(ch) else (False if LAT.match(ch) else None) for ch in s]
    nxt = None
    for i in range(len(kinds) - 1, -1, -1):
        if kinds[i] is None:
            kinds[i] = nxt
        else:
            nxt = kinds[i]
    prev = None
    for i, k in enumerate(kinds):
        if k is None:
            kinds[i] = prev
        prev = kinds[i] if kinds[i] is not None else prev
    runs = []
    for ch, k in zip(s, kinds):
        k = bool(k)
        if runs and runs[-1][1] == k:
            runs[-1][0] += ch
        else:
            runs.append([ch, k])
    return runs


def apply_font(cell):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    v = cell.value
    if v is None:
        return
    f = copy(cell.font)
    if isinstance(v, str) and v.startswith("="):
        f.name = EN
        cell.font = f
        return
    if not isinstance(v, str):
        f.name = EN
        cell.font = f
        return
    has_kr = bool(HAN.search(v))
    f.name = KR if has_kr else EN
    cell.font = f
    runs = split_runs(v)
    if has_kr and len(runs) > 1:
        cell.value = CellRichText([TextBlock(InlineFont(rFont=KR if k else EN, sz=f.sz, b=f.b, i=f.i, u=f.u,
                                   color=copy(f.color) if f.color else None), t) for t, k in runs])


def write_xlsx(path, corp, basis, reports, results):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter as L
    BLUE, GREY = "0000FF", "808080"
    HDR, WARN = PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="FFF2CC")
    NUM, PCT = '#,##0;(#,##0);"-"', '0.0%;-0.0%;"-"'

    def put(ws, r, c, v, color="000000", bold=False, fmt=None, fill=None, align=None):
        x = ws.cell(r, c, v)
        x.font = Font(size=10, color=color, bold=bold)
        if fmt:
            x.number_format = fmt
        if fill:
            x.fill = fill
        if align:
            x.alignment = Alignment(horizontal=align)
        return x

    pers = [r["period"] for r in reports if r["rcept_no"] in results]
    # 계정 순서: 최신 보고서 순서 우선, 과거에만 있는 계정은 뒤에 추가
    accts = []
    for p in reversed(pers):
        rn = next(r["rcept_no"] for r in reports if r["period"] == p)
        for a in results[rn]["accts"]:
            if a not in accts:
                accts.append(a)
    by_p = {r["period"]: results[r["rcept_no"]] for r in reports if r["rcept_no"] in results}
    mul = {p: TO_MN.get(by_p[p]["unit"], 1e-3) for p in pers}

    wb = Workbook()
    ws = wb.active
    ws.title = "판관비(분기)"
    wc = wb.create_sheet("누적(입력)")
    w3 = wb.create_sheet("3개월(입력)")
    wl = wb.create_sheet("목록")

    # 입력 시트: 행 = 계정 + 공시 합계, 열 = 보고서 기간
    R0 = 4
    in_row = {a: R0 + i for i, a in enumerate(accts)}
    tot_in = R0 + len(accts)
    p_col = {p: 3 + i for i, p in enumerate(pers)}
    for wsx, key, title in ((wc, 1, "누적 금액"), (w3, 0, "3개월 금액 (1분기는 누적과 같음, 사업보고서는 3개월 미공시)")):
        put(wsx, 1, 2, f"{corp} 판매비와관리비 {title} — {basis}, 백만원, DART 정기보고서 원문", bold=True)
        put(wsx, R0 - 1, 2, "계정", bold=True, fill=HDR)
        for p, c in p_col.items():
            put(wsx, R0 - 1, c, p, bold=True, fill=HDR, align="center")
        for a, r in in_row.items():
            put(wsx, r, 2, DISPLAY.get(a, a))
        put(wsx, tot_in, 2, "판매비와관리비(공시 합계)", bold=True)
        for p, c in p_col.items():
            d = by_p[p]
            if key == 0 and p.endswith(".12"):
                continue
            for a, (m3, cum) in d["accts"].items():
                v = (m3, cum)[key]
                if v is not None:
                    put(wsx, in_row[a], c, round(v * mul[p], 6), color=BLUE, fmt=NUM)
            if d["total"] is not None and d["total"][key] is not None:
                put(wsx, tot_in, c, round(d["total"][key] * mul[p], 6), color=BLUE, fmt=NUM, bold=True)
        wsx.column_dimensions["B"].width = 26
        for c in p_col.values():
            wsx.column_dimensions[L(c)].width = 11

    # 모델 시트
    years = sorted({int(p[:4]) for p in pers})
    qcols = {(y, q): 4 + 4 * i + (q - 1) for i, y in enumerate(years) for q in (1, 2, 3, 4)}
    ycol = {y: 4 + 4 * len(years) + i for i, y in enumerate(years)}
    ncol = 4 + 5 * len(years)
    put(ws, 1, 2, f"{corp} 판매비와관리비 분기 실적 ({basis}, 백만원)", bold=True)
    put(ws, 2, 2, "1Q = 1분기 누적, 2Q·3Q = 해당 보고서 3개월, 4Q = 연간 − 3분기 누적. 출처: DART 분기·반기·사업보고서 주석", color=GREY)
    H = 4
    put(ws, H, 2, "계정", bold=True, fill=HDR)
    put(ws, H, 3, "구분", bold=True, fill=HDR, align="center")
    for (y, q), c in qcols.items():
        put(ws, H, c, f"{q}Q{str(y)[2:]}", bold=True, fill=HDR, align="center")
    for y, c in ycol.items():
        put(ws, H, c, str(y), bold=True, fill=HDR, align="center")
    put(ws, H, ncol, "분류 근거", bold=True, fill=HDR)

    def src(sheet, period, row):
        return f"'{sheet}'!{L(p_col[period])}{row}"

    def qformula(y, q, row):
        """분기 셀 수식 (입력 시트 참조). 해당 기간 보고서가 없으면 None"""
        if q < 4:
            p = f"{y}.{q * 3:02d}"
            return f'=IFERROR({src("3개월(입력)", p, row)},"n/a")' if p in p_col else None
        p12, p09 = f"{y}.12", f"{y}.09"
        if p12 in p_col and p09 in p_col:
            return f'=IFERROR({src("누적(입력)", p12, row)}-{src("누적(입력)", p09, row)},"n/a")'
        return None

    r1 = H + 1
    for i, a in enumerate(accts):
        r = r1 + i
        kind, why, unknown = classify(a)
        put(ws, r, 2, DISPLAY.get(a, a))
        put(ws, r, 3, kind, align="center", fill=WARN if unknown else None)
        for (y, q), c in qcols.items():
            f = qformula(y, q, in_row[a])
            if f:
                put(ws, r, c, f, fmt=NUM, align="right")
        for y, c in ycol.items():
            q1, q4 = L(qcols[(y, 1)]), L(qcols[(y, 4)])
            put(ws, r, c, f'=IFERROR(SUM({q1}{r}:{q4}{r}),"n/a")', fmt=NUM, align="right")
        put(ws, r, ncol, why, color=GREY)
    rn = r1 + len(accts) - 1
    rs, rp, rk = rn + 1, rn + 2, rn + 3
    put(ws, rs, 2, "판매비와관리비 합계", bold=True)
    put(ws, rp, 2, "공시 합계", color=GREY)
    put(ws, rk, 2, "Check (합계 − 공시)", color=GREY)
    for (y, q), c in list(qcols.items()) + [((y, 0), c) for y, c in ycol.items()]:
        C = L(c)
        put(ws, rs, c, f'=IFERROR(SUM({C}{r1}:{C}{rn}),"n/a")', bold=True, fmt=NUM, align="right")
        if q == 0:
            fp = f'=IFERROR({src("누적(입력)", f"{y}.12", tot_in)},"n/a")' if f"{y}.12" in p_col else None
        else:
            fp = qformula(y, q, tot_in)
        if fp:
            put(ws, rp, c, fp, color=GREY, fmt=NUM, align="right")
            put(ws, rk, c, f'=IFERROR({C}{rs}-{C}{rp},"n/a")', color=GREY, fmt=NUM, align="right")
    rf, rv, rr = rk + 2, rk + 3, rk + 4
    put(ws, rf, 2, "고정비")
    put(ws, rv, 2, "변동비")
    put(ws, rr, 2, "고정비 비중")
    for c in list(qcols.values()) + list(ycol.values()):
        C = L(c)
        put(ws, rf, c, f'=IFERROR(SUMIF($C${r1}:$C${rn},"고정비",{C}{r1}:{C}{rn}),"n/a")', fmt=NUM, align="right")
        put(ws, rv, c, f'=IFERROR(SUMIF($C${r1}:$C${rn},"변동비",{C}{r1}:{C}{rn}),"n/a")', fmt=NUM, align="right")
        put(ws, rr, c, f'=IFERROR({C}{rf}/{C}{rs},"n/a")', fmt=PCT, align="right")
    put(ws, rr + 2, 2, "노란 칸 = 분류표에 없는 계정(고정비 기본값) — 확인 후 수정. 4Q는 사업보고서 연간에서 3분기 누적을 빼서 계산하므로 "
        "연간 보고서에서 과거 분기가 재작성되면 그 차이가 4Q에 들어감", color=GREY)
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 8
    for c in list(qcols.values()) + list(ycol.values()):
        ws.column_dimensions[L(c)].width = 9
    ws.column_dimensions[L(ncol)].width = 30

    # 목록
    put(wl, 1, 1, f"{corp} · {basis} · 금액은 백만원 환산", bold=True)
    for j, h in enumerate(["기간", "보고서", "접수번호", "접수일", "원문 단위", "판관비 표", "표 위 제목", "DART 링크"], 1):
        put(wl, 2, j, h, bold=True, fill=HDR)
    for i, r in enumerate(reports, 3):
        d = results.get(r["rcept_no"])
        row = [r["period"], r["report_nm"], r["rcept_no"], r.get("rcept_dt", ""), d["unit"] if d else "",
               "Y" if d else "못 찾음", d["context"] if d else "", f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={r['rcept_no']}"]
        for j, v in enumerate(row, 1):
            put(wl, i, j, v, fill=None if d else WARN)
    for c, w in zip("ABCDEFGH", (9, 30, 16, 10, 9, 9, 30, 60)):
        wl.column_dimensions[c].width = w

    for w in wb.worksheets:
        for row in w.iter_rows():
            for x in row:
                apply_font(x)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("stock", help="종목코드 6자리 (엘앤에프 066970)")
    ap.add_argument("--basis", default="연결", choices=["연결", "별도"])
    ap.add_argument("--since", default="20190101")
    ap.add_argument("--until", default=time.strftime("%Y%m%d"))
    ap.add_argument("--out", default=None)
    ap.add_argument("--env", default=None)
    ap.add_argument("--keep-xml", default=None)
    ap.add_argument("--xml-dir", default=None, help="저장해 둔 원문 XML 폴더 ({기간}_{접수번호}.xml) — 네트워크 없이 실행")
    a = ap.parse_args()

    if a.xml_dir:
        files = sorted(Path(a.xml_dir).glob("*.xml"))
        name = a.stock
        reps = []
        for f in files:
            m = re.match(r"(\d{4}\.\d{2})_(\d+)", f.stem)
            if m and a.since[:4] <= m.group(1)[:4]:
                kind = {"03": "분기보고서", "06": "반기보고서", "09": "분기보고서", "12": "사업보고서"}[m.group(1)[5:]]
                reps.append({"period": m.group(1), "kind": kind, "rcept_no": m.group(2),
                             "report_nm": f"{kind} ({m.group(1)})", "rcept_dt": "", "_file": f})
        get = lambda r: r["_file"].read_text(encoding="utf-8")  # noqa: E731
    else:
        key = load_key(a.env)
        cc, name = corp_code(key, a.stock)
        reps = periodic_reports(key, cc, a.since, a.until)
        get = lambda r: fetch_doc(key, r["rcept_no"])  # noqa: E731
    print(f"{name}: 정기보고서 {len(reps)}건")
    results = {}
    for r in reps:
        xml = get(r)
        if xml is None:
            print(f"  {r['period']} 원문 없음")
            continue
        if a.keep_xml and not a.xml_dir:
            Path(a.keep_xml).mkdir(parents=True, exist_ok=True)
            Path(a.keep_xml, f"{r['period']}_{r['rcept_no']}.xml").write_text(xml, encoding="utf-8")
        d = extract_sga(xml, a.basis)
        if d:
            results[r["rcept_no"]] = d
            print(f"  {r['period']} {r['kind']}: 계정 {len(d['accts'])}개, 단위 {d['unit']}")
        else:
            print(f"  {r['period']} {r['kind']}: 판관비 표를 찾지 못함")
        if not a.xml_dir:
            time.sleep(0.4)
    out = a.out or f"{a.stock}_sga_{a.basis}.xlsx"
    write_xlsx(out, name, a.basis, reps, results)
    print(f"완료 → {out}")


if __name__ == "__main__":
    main()
